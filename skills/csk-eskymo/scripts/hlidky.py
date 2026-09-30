# -*- coding: utf-8 -*-
"""
hlidky — startovka závodu družstev (hlídek) v Eskymu.

Kroky (každý je samostatný, mezivýstupy jsou JSON k nahlédnutí):
  1) python hlidky.py loni vysledky_loni.ods -o loni.json
       loňské pořadí hlídek z Eskymo sešitu (.ods) nebo z Export XLS (.xlsx);
       pořadí = sloupec poř., DNF/DSQ na konci v pořadí řádků.
  2) python hlidky.py nasad hlidky.json zavod.ods [--loni loni.json] -c nastaveni.json -o plan.json
       S26: MČR družstev „v obráceném pořadí loňských výsledků“ (vítěz startuje poslední).
       Návaznost letošní hlídky na loňskou: (a) identita oddíl+písmeno, když má letos
       hlídka písmeno deklarované; (b) shoda přes závodníka, ale jen v TÉMŽE oddílu;
       (c) zbylá loňská umístění oddílu podle síly (VT); (d) jinak „nová“.
       U každé hlídky vypíše, podle čeho byla nasazena — to kontroluje pořadatel.
  3) python hlidky.py zapis plan.json zavod.ods vystup.ods [-c nastaveni.json]
       přidělí čísla (cisla.py), zapíše list `hlidky` (A–F) a `<kat>_sl` (B stč, P H-RGC)
       ve fyzickém pořadí startu; na kopii opraví překlep `uuper(` → `UPPER(`.

hlidky.json — seznam hlídek (oddíly často přihlásí jen první loď, zbytek dodají později):
  [{"kat": "k1m", "lode": ["9162", "9091", "9184"], "pismeno": "A", "oddil": "USK Pha",
    "poradi": 3, "zdroj": "e-mail USK"}]
  C2: loď = "57036 57054". Chybějící lodě vynech (lode: ["9162"]). `oddil` jinak z reg.
  Z CSV exportu: python hlidky.py z-prihlasek lode.json -o hlidky.json (1 řádek = 1 hlídka).

nastaveni.json:
  {"poradi_kategorii": ["c1m","k1z","k1m","c1z","c2m"],
   "nasazeni": {"nove": "zacatek", "nove_razeni": "vt-nejslabsi-prvni", "pismena": "podle-nasazeni"},
   "oddil_alias": {"Vys.Mýto": "SKK VM"},
   "cisla": {"rezim": "sestupne", "start": 70, "mezera": 5, "vynechat": [41, 66]}}
  pismena: deklarovana | podle-nasazeni (A = nejlépe nasazená) | podle-prihlasky | zadna
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict

import cisla as C
from domena import kategorie_eskymo, lepsi_vt, lod as typ_lode, vt_rank
from eskymo_ods import Workbook, norm_rgc, rgc_cell_value
from inspect_workbook import read_param
from registr import Registr

PISMENA = 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'


# ---------- 1) loňské výsledky ----------
def _split(v) -> list[str]:
    return [x.strip() for x in str(v or '').split('\n') if x.strip()]


def _klub_a_pismeno(oddil_bunka) -> tuple[str, str]:
    radky = _split(oddil_bunka)
    if not radky:
        return '', ''
    prvni = radky[0]
    posl = radky[-1].split()
    pism = posl[-1] if posl and len(posl[-1]) == 1 and posl[-1] in PISMENA else ''
    if pism and prvni.endswith(' ' + pism):
        prvni = prvni[:-2].strip()
    return prvni, pism


def nacti_loni(path: str) -> dict:
    """{kat: [{'poradi': n, 'umisteni': '3.'|'DNF', 'oddil': .., 'pismeno': .., 'rgc': [..]}]}"""
    out = {}
    if path.lower().endswith('.ods'):
        wb = Workbook(path)
        param = read_param(wb) if wb.has('param') else {'kategorie': []}
        kats = [k['kod'].lower() for k in param['kategorie']] or \
               [n for n in wb.sheet_names() if len(n) == 3 and wb.has(n + '_sl')]
        for kat in kats:
            if not wb.has(kat):
                continue
            rows = wb.sheet(kat).values(max_cols=12)[2:]
            # sloupce výsledkového listu: C poř. | G rgc | K oddíl
            data = [(r[2] if len(r) > 2 else None, r[6] if len(r) > 6 else None, r[10] if len(r) > 10 else None)
                    for r in rows]
            out[kat] = _poradi(data)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        for kat in wb.sheetnames:
            if kat == 'param':
                continue
            rows = list(wb[kat].iter_rows(values_only=True))[2:]
            # Export XLS: A poř. | E rgc | I oddíl
            out[kategorie_eskymo(kat) or kat] = _poradi([(r[0], r[4], r[8]) for r in rows if len(r) > 8])
    return out


def _poradi(data) -> list[dict]:
    rows = []
    for i, (por, rgc, oddil) in enumerate(data):
        rgcs = [norm_rgc(x) for x in _split(rgc)]
        rgcs = [x for x in rgcs if x and x != '0']
        if not rgcs:
            continue
        klub, pism = _klub_a_pismeno(oddil)
        s = str(por or '').strip().rstrip('.')
        rows.append({'fyz': i, 'umisteni': (s + '.') if s.isdigit() else 'DNF/DSQ',
                     'num': int(s) if s.isdigit() else None, 'oddil': klub, 'pismeno': pism,
                     'rgc': sorted({p for b in rgcs for p in b.split()})})
    rows.sort(key=lambda t: (t['num'] is None, t['num'] or 0, t['fyz']))
    for n, t in enumerate(rows, 1):
        t['poradi'] = n
        del t['fyz'], t['num']
    return rows


# ---------- 2) nasazení ----------
def z_prihlasek(lode: list[dict]) -> list[dict]:
    """lode.json z prihlasky.py (1 řádek exportu = 1 hlídka, zatím jen první loď)."""
    return [{'kat': l['kat'], 'lode': [' '.join(l['rgc'])], 'poradi': l.get('poradi'),
             'zdroj': l.get('zdroj'), 'pismeno': l.get('pismeno', '')} for l in lode]


def nasad(hlidky: list[dict], wb: Workbook, loni: dict | None, nast: dict) -> dict:
    reg = Registr.ze_sesitu(wb)
    pole = {k: v['hodnota'] for k, v in read_param(wb)['pole'].items()}
    disc = str(pole.get('Disciplína') or 'slalom')
    alias = nast.get('oddil_alias', {})
    n = nast.get('nasazeni', {})
    por_kat = [kategorie_eskymo(k) or k for k in nast.get('poradi_kategorii') or
               [k['kod'] for k in read_param(wb)['kategorie']]]
    nezname = [h for h in hlidky if not kategorie_eskymo(h.get('kat', ''))]
    if nezname:
        raise SystemExit(f"Hlídky s neznámou kategorií: {[(h.get('kat'), h.get('zdroj')) for h in nezname]}")
    chybi = sorted({kategorie_eskymo(h['kat']) for h in hlidky} - set(por_kat))
    if chybi:
        raise SystemExit(f'Kategorie {chybi} mají přihlášené hlídky, ale nejsou v poradi_kategorii — doplň je '
                         '(jinak by hlídky tiše vypadly ze startovky).')
    plan = {'zavod': {k: pole.get(k) for k in ('Název závodu', 'Datum závodu', 'Číslo závodu', 'Disciplína')},
            'nastaveni': nast, 'kategorie': [], 'nesparovana_loni': {}, 'k_rozhodnuti': []}

    for i, h in enumerate(hlidky):
        h['kat'] = kategorie_eskymo(h['kat']) or h['kat']
        h['lode'] = [norm_rgc(x) for x in h.get('lode', []) if norm_rgc(x)]
        prvni = reg.get(h['lode'][0].split()[0]) if h['lode'] else None
        h['oddil'] = h.get('oddil') or (prvni.oddil if prvni else '')
        h['vt'] = lepsi_vt(*[reg.vt_lode(b.split(), h['kat'], disc) for b in h['lode']]) if h['lode'] else ''
        h.setdefault('poradi', i)
        h['jmena'] = [' / '.join(reg.get(p).cele_jmeno for p in b.split() if reg.get(p)) for b in h['lode']]

    for kat in por_kat:
        ents = [h for h in hlidky if h['kat'] == kat]
        if not ents:
            continue
        teams = [dict(t, oddil=alias.get(t['oddil'], t['oddil'])) for t in (loni or {}).get(kat, [])]
        obs = set()
        # (a) identita oddíl + deklarované písmeno
        for e in ents:
            if e.get('pismeno'):
                for t in teams:
                    if t['poradi'] not in obs and t['oddil'] == e['oddil'] and t['pismeno'] == e['pismeno']:
                        e.update(loni=t['poradi'], loni_umisteni=t['umisteni'],
                                 vazba=f"identita {e['oddil']} {e['pismeno']}")
                        obs.add(t['poradi'])
                        break
        # (b) přes závodníka v témže oddílu
        for e in ents:
            if 'loni' in e:
                continue
            osoby = {p for b in e['lode'] for p in b.split()}
            for t in teams:
                if t['poradi'] in obs or t['oddil'] != e['oddil']:
                    continue
                spol = osoby & set(t['rgc'])
                if spol:
                    o = reg.get(sorted(spol)[0])
                    e.update(loni=t['poradi'], loni_umisteni=t['umisteni'],
                             vazba=f"osoba {o.cele_jmeno if o else sorted(spol)[0]}")
                    obs.add(t['poradi'])
                    break
        # (c) zbylá umístění oddílu podle síly (VT, pak pořadí v přihláškách)
        for e in sorted([e for e in ents if 'loni' not in e], key=lambda e: (vt_rank(e['vt']), e['poradi'])):
            for t in teams:
                if t['poradi'] not in obs and t['oddil'] == e['oddil']:
                    e.update(loni=t['poradi'], loni_umisteni=t['umisteni'], vazba='oddíl (zbylé umístění)')
                    obs.add(t['poradi'])
                    plan['k_rozhodnuti'].append(
                        f"{kat.upper()}: hlídka {e['oddil']} ({'; '.join(e.get('jmena') or e['lode'])}) zdědila "
                        f"loňské {t['umisteni']} místo jen podle oddílu (žádný společný závodník) — potvrdit.")
                    break
        for e in ents:
            e.setdefault('loni', None)
            e.setdefault('loni_umisteni', None)
            if 'vazba' not in e:
                # sdílí člena s loňskou hlídkou, jejíž umístění už zdědila jiná letošní hlídka oddílu
                osoby = {p for b in e['lode'] for p in b.split()}
                sdil = [t for t in teams if t['oddil'] == e['oddil'] and osoby & set(t['rgc'])]
                if sdil:
                    t = sdil[0]
                    e['vazba'] = f"jako nová — loni {t['umisteni']} {t['oddil']} zdědila jiná hlídka oddílu"
                    plan['k_rozhodnuti'].append(
                        f"{kat.upper()}: dvě letošní hlídky oddílu {e['oddil']} navazují na tutéž loňskou "
                        f"({t['umisteni']}); umístění zdědila jen jedna — potvrdit, která.")
                else:
                    e['vazba'] = 'nová (loni bez návaznosti)'
        nesp = [t for t in teams if t['poradi'] not in obs]
        if nesp:
            plan['nesparovana_loni'][kat] = [f"{t['umisteni']} {t['oddil']} {t['pismeno']}".strip() for t in nesp]

        nove = [e for e in ents if e['loni'] is None]
        nejslabsi = n.get('nove_razeni', 'vt-nejslabsi-prvni') == 'vt-nejslabsi-prvni'
        nove.sort(key=lambda e: ((-vt_rank(e['vt'])) if nejslabsi else vt_rank(e['vt']), e['poradi']))
        stare = sorted([e for e in ents if e['loni'] is not None], key=lambda e: -e['loni'])
        poradi = (nove + stare) if n.get('nove', 'zacatek') == 'zacatek' else (stare + nove)

        rezim = n.get('pismena', 'podle-nasazeni')
        if rezim != 'deklarovana':
            po_odd = defaultdict(list)
            zdroj = list(reversed(poradi)) if rezim == 'podle-nasazeni' else sorted(poradi, key=lambda e: e['poradi'])
            for e in zdroj:
                po_odd[e['oddil']].append(e)
            for lst in po_odd.values():
                for i, e in enumerate(lst):
                    e['pismeno'] = PISMENA[i] if len(lst) > 1 and rezim != 'zadna' else ''
        for i, e in enumerate(poradi, 1):
            e['hrgc'] = f'{kat.upper()}-{i:02d}'
        plan['kategorie'].append({'kat': kat, 'hlidky': poradi})
    plan['k_rozhodnuti'] += kontrola_clenu(plan, reg)
    return plan


def kontrola_clenu(plan: dict, reg) -> list[str]:
    """P 2.09.02 / 2.39.03: žena v mužské i ženské hlídce téže lodi, víc než 3 hlídky, 2× v kategorii."""
    kde = defaultdict(list)
    for k in plan['kategorie']:
        for h in k['hlidky']:
            for b in h['lode']:
                for p in b.split():
                    kde[p].append(k['kat'])
    out = []
    for p, kats in kde.items():
        o = reg.get(p)
        jm = f'{o.cele_jmeno} ({p})' if o else p
        for k in sorted(set(kats)):
            if kats.count(k) > 1:
                out.append(f'{jm} je ve {kats.count(k)} hlídkách kategorie {k.upper()} — v kategorii smí jen jednou.')
        if len(kats) > 3:
            out.append(f'{jm} je v {len(kats)} hlídkách — na MČR/ČP/NKZ max. 3 družstva (P 2.39.03).')
        for lod in ('k1', 'c1', 'c2'):
            if f'{lod}m' in kats and f'{lod}z' in kats:
                out.append(f'{jm} je v mužské i ženské hlídce {lod.upper()} — žena smí v mužském družstvu jen když '
                           'nejede v ženském družstvu téže lodní kategorie (P 2.09.02).')
    return out


# ---------- 3) zápis ----------
def zapis(plan: dict, wb: Workbook, nast: dict, prepsat=False) -> dict:
    if not wb.has('hlidky'):
        raise SystemExit('Sešit nemá list `hlidky` — v Eskymu musí být při Nový závod zvoleno Hlídky: ano.')
    oprav = wb.patch_formulas('uuper(', 'UPPER(')
    kats = [k for k in plan['kategorie'] if k['hlidky']]
    cn = {**{'rezim': 'desitky', 'start': 1, 'mezera': 0, 'vynechat': []},
          **plan.get('nastaveni', {}).get('cisla', {}), **(nast or {}).get('cisla', {})}
    pevne = {k['kat']: [h.get('stc') for h in k['hlidky']] for k in kats} if cn['rezim'] == 'pevne' else None
    res = C.prirad([(k['kat'], len(k['hlidky'])) for k in kats], cn['rezim'], cn.get('start'),
                   int(cn.get('mezera') or 0), cn.get('vynechat') or [], int(cn.get('zarovnani') or 10),
                   cn.get('max'), pevne)
    hs = wb.sheet('hlidky')
    hvals = hs.values(max_cols=6)
    volne = [r for r in range(1, max(len(hvals), 2) + 400)
             if r >= len(hvals) or all(v in (None, '') for v in (hvals[r] + [None] * 6)[:6])]
    obsazene = [r for r in range(1, len(hvals)) if hvals[r] and any(v not in (None, '') for v in hvals[r][:6])]
    if obsazene and not prepsat:
        raise SystemExit(f'List hlidky už obsahuje {len(obsazene)} hlídek — použij čistou šablonu nebo --prepsat.')
    if prepsat:
        for r in obsazene:
            for c in range(6):
                hs.set_input(r, c, None)
        volne = sorted(set(volne) | set(obsazene))
    vsechny = [h for k in kats for h in k['hlidky']]
    for r, h in zip(volne, vsechny):
        hs.set_input(r, 0, h['hrgc'])
        hs.set_input(r, 1, typ_lode(h['kat']))
        for j in range(3):
            hs.set_input(r, 2 + j, rgc_cell_value(h['lode'][j]) if j < len(h['lode']) else None)
        hs.set_input(r, 5, h.get('pismeno') or None)
    for k in kats:
        sh = wb.sheet(f"{k['kat']}_sl")
        vals = sh.values(max_cols=16)
        rows = [r for r in range(2, len(vals)) if vals[r] and isinstance(vals[r][0], float)]
        if len(k['hlidky']) > len(rows):
            raise SystemExit(f"{k['kat']}_sl má jen {len(rows)} řádků a hlídek je {len(k['hlidky'])}")
        for r in rows:
            v = (vals[r] + [None] * 16)
            if (v[1] not in (None, '') or v[15] not in (None, '')):
                if not prepsat:
                    raise SystemExit(f"{k['kat']}_sl už obsahuje data — čistá šablona nebo --prepsat")
                sh.set_input(r, 1, None)
                sh.set_input(r, 15, None)
        for i, h in enumerate(k['hlidky']):
            h['stc'] = res['cisla'][k['kat']][i]
            sh.set_input(rows[i], 1, h['stc'])
            sh.set_input(rows[i], 15, h['hrgc'])
    return {'cisla': res, 'uuper': oprav}


def prehled(plan: dict) -> str:
    L = [f"# Nasazení hlídek — {plan['zavod'].get('Název závodu')} ({plan['zavod'].get('Datum závodu')})", '']
    for k in plan['kategorie']:
        L += [f"## {k['kat'].upper()} ({len(k['hlidky'])} hlídek)",
              '| start | stč | H-RGC | oddíl | lodě | VT | loni | nasazeno podle |', '|---:|---:|---|---|---|---|---|---|']
        for i, h in enumerate(k['hlidky'], 1):
            L.append(f"| {i} | {h.get('stc') or ''} | {h['hrgc']} | {h['oddil']} {h.get('pismeno') or ''} | "
                     f"{'; '.join(h.get('jmena') or h['lode'])} | {h['vt'] or '—'} | {h['loni_umisteni'] or '—'} | "
                     f"{h['vazba']} |")
        if plan['nesparovana_loni'].get(k['kat']):
            L.append(f"\n! Loňské hlídky bez letošní návaznosti: {', '.join(plan['nesparovana_loni'][k['kat']])} "
                     '— zkontroluj přejmenované oddíly (oddil_alias) a přestupy.')
        L.append('')
    if plan.get('k_rozhodnuti'):
        L += ['## K rozhodnutí pořadatele', *(f'- {x}' for x in dict.fromkeys(plan['k_rozhodnuti'])), '']
    if plan.get('opravy_sablony'):
        L += ['## Opravy šablony (jen ve výstupní kopii)', *(f'- {x}' for x in plan['opravy_sablony']), '']
    return '\n'.join(L)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    a0 = sub.add_parser('loni'); a0.add_argument('soubor'); a0.add_argument('-o', '--out', required=True)
    a1 = sub.add_parser('z-prihlasek'); a1.add_argument('lode'); a1.add_argument('-o', '--out', required=True)
    a2 = sub.add_parser('nasad'); a2.add_argument('hlidky'); a2.add_argument('sesit')
    a2.add_argument('--loni'); a2.add_argument('-c', '--nastaveni'); a2.add_argument('-o', '--out', required=True)
    a2.add_argument('--prehled')
    a3 = sub.add_parser('zapis'); a3.add_argument('plan'); a3.add_argument('sesit'); a3.add_argument('vystup')
    a3.add_argument('-c', '--nastaveni'); a3.add_argument('--prepsat', action='store_true'); a3.add_argument('--prehled')
    a = ap.parse_args(argv)
    nast = json.load(open(a.nastaveni, encoding='utf-8')) if getattr(a, 'nastaveni', None) else {}

    if a.cmd == 'loni':
        d = nacti_loni(a.soubor)
        json.dump(d, open(a.out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        print(', '.join(f'{k}: {len(v)}' for k, v in d.items()))
        return 0
    if a.cmd == 'z-prihlasek':
        d = json.load(open(a.lode, encoding='utf-8'))
        d = d['lode'] if isinstance(d, dict) else d
        json.dump(z_prihlasek(d), open(a.out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        return 0
    if a.cmd == 'nasad':
        hl = json.load(open(a.hlidky, encoding='utf-8'))
        loni = json.load(open(a.loni, encoding='utf-8')) if a.loni else None
        plan = nasad(hl, Workbook(a.sesit), loni, nast)
        json.dump(plan, open(a.out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        t = prehled(plan)
        if a.prehled:
            open(a.prehled, 'w', encoding='utf-8').write(t)
        print(t)
        return 0
    plan = json.load(open(a.plan, encoding='utf-8'))
    wb = Workbook(a.sesit)
    try:
        res = zapis(plan, wb, nast, a.prepsat)
    except C.ChybaCisel as e:
        print(f'ČÍSLOVÁNÍ: {e}', file=sys.stderr)
        return 1
    if res['uuper']:
        plan['opravy_sablony'] = [f"překlep uuper( → UPPER( ve vzorcích listu hlidky ({res['uuper']} buněk; "
                                  'jinak #NAME? ve VT u C2 hlídek) — nahlásit autorovi Eskyma']
    wb.save(a.vystup)
    json.dump(plan, open(a.plan, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    if res['uuper']:
        print(f"Opraven překlep šablony uuper( → UPPER( v {res['uuper']} buňkách (jen v této kopii).")
    print(C.popis(res['cisla'], [k['kat'] for k in plan['kategorie']]))
    if a.prehled:
        open(a.prehled, 'w', encoding='utf-8').write(prehled(plan))
    print(f'Zapsáno → {a.vystup}. Ověř: verify_workbook.py {a.vystup}')
    return 0


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    sys.exit(main())
