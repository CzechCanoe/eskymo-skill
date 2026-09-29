# -*- coding: utf-8 -*-
"""
check_rules — platí pro danou sezónu pravidla a Směrnice, ze kterých skill vychází?

    python check_rules.py [--sezona 2027] [--sesit zavod.ods] [--stahnout DIR]

Skill má zadrátovaný výtah (references/pravidla-<rok>.md) a manifest zdrojů
(references/pravidla-manifest.json: URL, velikosti, sha256, Last-Modified).
Tenhle skript na kanoe.cz ověří:
  1. jestli pro sezónu závodu existují novější Směrnice, než ze kterých je výtah,
  2. jestli se známé dokumenty nezměnily (tiché opravy příloh: „-oprava“, jiná velikost/datum),
  3. jestli nevyšlo nové vydání Pravidel.
Sezóna = --sezona, jinak rok z param „Datum závodu“ (--sesit), jinak letošní rok.

Výstup: verdikt + co dělat. Návratový kód 0 = výtah platí, 3 = nutná aktualizace
nebo kontrola změn, 2 = nelze ověřit (offline) → pracuj s výtahem a řekni to pořadateli.
--stahnout DIR stáhne aktuální Směrnice a přílohy (docx převede na .txt), aby je
agent mohl přečíst a porovnat s výtahem.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import sys
import urllib.error
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
MANIFEST = os.path.join(HERE, '..', 'references', 'pravidla-manifest.json')
BASE = 'https://www.kanoe.cz'


def _req(url, method='GET'):
    m = json.load(open(MANIFEST, encoding='utf-8'))
    ua = m['detection']['user_agent']
    return urllib.request.Request(url, method=method, headers={'User-Agent': ua})


def get(url) -> str:
    with urllib.request.urlopen(_req(url), timeout=60) as r:
        return r.read().decode('utf-8', errors='replace')


def head(url) -> dict:
    try:
        with urllib.request.urlopen(_req(url, 'HEAD'), timeout=60) as r:
            return {'status': r.status, 'bytes': int(r.headers.get('Content-Length') or 0) or None,
                    'last_modified': r.headers.get('Last-Modified')}
    except urllib.error.HTTPError as e:
        return {'status': e.code}


def _sezona_ze_sesitu(path):
    sys.path.insert(0, HERE)
    from eskymo_ods import Workbook
    from inspect_workbook import read_param
    d = read_param(Workbook(path))['pole'].get('Datum závodu', {}).get('hodnota')
    m = re.search(r'(\d{2})\s*$', str(d or ''))
    return 2000 + int(m.group(1)) if m else None


def zkontroluj(sezona: int) -> dict:
    man = json.load(open(MANIFEST, encoding='utf-8'))
    det = man['detection']
    s_man = int(man['current']['smernice_season'])
    rep = {'sezona': sezona, 'vytah_pro': s_man, 'zjisteni': [], 'verdikt': 'OK', 'odkazy': {}}

    listing = get(det['listing_url'])
    roky = {}
    for href, rok in re.findall(det['smernice_article_regex'], listing):
        if not re.search(det['smernice_exclude_regex'], href):
            roky.setdefault(int(rok), BASE + href)
    nejnovejsi = max(roky) if roky else None
    rep['nejnovejsi_smernice_na_webu'] = nejnovejsi

    if sezona > s_man:
        if sezona in roky:
            rep['verdikt'] = 'NOVA_SEZONA'
            rep['odkazy']['smernice'] = roky[sezona]
            rep['zjisteni'].append(f'Pro sezónu {sezona} jsou zveřejněné nové Směrnice ({roky[sezona]}), '
                                   f'výtah skillu je pro {s_man}.')
        else:
            rep['verdikt'] = 'NEZVEREJNENO'
            rep['zjisteni'].append(f'Směrnice {sezona} zatím nejsou zveřejněné (obvykle prosinec–únor). '
                                   f'Platí poslední schválené — výtah {s_man} — s výhradou.')
    elif sezona < s_man:
        rep['verdikt'] = 'STARSI_SEZONA'
        rep['zjisteni'].append(f'Závod je ze sezóny {sezona}; výtah je pro {s_man}. Pro zpětné zpracování '
                               f'použij Směrnice {sezona}: {roky.get(sezona, det["smernice_article_url_template"].replace("{YYYY}", str(sezona)))}')

    # změny známých dokumentů aktuální sezóny výtahu
    for d in man['documents']:
        if d.get('status') != 'current':
            continue
        h = head(d['url'])
        if h.get('status') != 200:
            rep['zjisteni'].append(f"{d['id']}: dokument není dostupný (HTTP {h.get('status')}) — mohl být nahrazen")
            rep['verdikt'] = rep['verdikt'] if rep['verdikt'] != 'OK' else 'ZMENA'
            continue
        zm = []
        if d.get('bytes') and h.get('bytes') and h['bytes'] != d['bytes']:
            zm.append(f"velikost {d['bytes']} → {h['bytes']}")
        if d.get('last_modified') and h.get('last_modified'):
            try:
                a = dt.datetime.strptime(h['last_modified'], '%a, %d %b %Y %H:%M:%S %Z')
                b = dt.datetime.fromisoformat(d['last_modified'].replace('Z', ''))
                if abs((a - b).total_seconds()) > 60:
                    zm.append(f"Last-Modified {d['last_modified']} → {h['last_modified']}")
            except ValueError:
                pass
        if zm:
            rep['zjisteni'].append(f"{d['id']}: změna ({', '.join(zm)}) — stáhni a porovnej s výtahem")
            if rep['verdikt'] == 'OK':
                rep['verdikt'] = 'ZMENA'

    # nové soubory v článku aktuální sezóny (opravy příloh)
    cur = next((d for d in man['documents'] if d['id'] == man['current']['smernice']), None)
    if cur and cur.get('page_url'):
        page = get(cur['page_url'])
        zname = {d['url'].split('/')[-1] for d in man['documents']}
        for m in re.finditer(det['priloha_file_regex'], page):
            fn = m.group(1).split('/')[-1]
            if fn not in zname:
                rep['zjisteni'].append(f'nový soubor v článku Směrnic {s_man}: {fn}')
                if rep['verdikt'] == 'OK':
                    rep['verdikt'] = 'ZMENA'
        rep['odkazy'].setdefault('smernice', cur['page_url'])

    # Pravidla
    pp = get(det['pravidla_article_url'])
    ed = sorted({int(y) for _, y in re.findall(det['pravidla_file_regex'], pp)})
    if ed and ed[-1] > 2022:
        rep['zjisteni'].append(f'nové vydání Pravidel {ed[-1]} — výtah je z Pravidel 2022')
        rep['verdikt'] = 'NOVA_PRAVIDLA'
    return rep


def stahni(page_url: str, outdir: str) -> list[str]:
    os.makedirs(outdir, exist_ok=True)
    page = get(page_url)
    files = sorted(set(re.findall(r'href="(/img/CSKDV/[^"]+\.(?:docx?|xlsx?|pdf))"', page)))
    out = []
    for f in files:
        dst = os.path.join(outdir, f.split('/')[-1])
        with urllib.request.urlopen(_req(BASE + f), timeout=120) as r, open(dst, 'wb') as w:
            w.write(r.read())
        out.append(dst)
        if dst.endswith('.docx'):
            try:
                import docx
                d = docx.Document(dst)
                lines = [p.text for p in d.paragraphs]
                for t in d.tables:
                    lines += [' | '.join(c.text.strip() for c in row.cells) for row in t.rows]
                open(dst[:-5] + '.txt', 'w', encoding='utf-8').write('\n'.join(lines))
                out.append(dst[:-5] + '.txt')
            except ImportError:
                pass
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--sezona', type=int)
    ap.add_argument('--sesit')
    ap.add_argument('--stahnout', metavar='DIR')
    ap.add_argument('--json', action='store_true')
    a = ap.parse_args(argv)
    sezona = a.sezona or (a.sesit and _sezona_ze_sesitu(a.sesit)) or dt.date.today().year
    try:
        rep = zkontroluj(int(sezona))
    except (urllib.error.URLError, OSError) as e:
        man = json.load(open(MANIFEST, encoding='utf-8'))
        print(f'NELZE OVĚŘIT (offline? {e}). Pracuj s výtahem Směrnic {man["current"]["smernice_season"]} '
              f'a Pravidel 2022 a řekni pořadateli, že platnost nebyla ověřena.')
        return 2
    if a.json:
        print(json.dumps(rep, ensure_ascii=False, indent=1))
    else:
        print(f"Sezóna závodu: {rep['sezona']}   výtah skillu: Směrnice {rep['vytah_pro']} + Pravidla 2022   "
              f"nejnovější Směrnice na webu: {rep.get('nejnovejsi_smernice_na_webu')}")
        print(f"VERDIKT: {rep['verdikt']}")
        for z in rep['zjisteni']:
            print(f'  - {z}')
        if rep['verdikt'] == 'OK':
            print('Výtah references/pravidla-2026.md platí.')
    if a.stahnout and rep['odkazy'].get('smernice'):
        for f in stahni(rep['odkazy']['smernice'], a.stahnout):
            print('staženo:', f)
    return 0 if rep['verdikt'] in ('OK', 'NEZVEREJNENO') else 3


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    sys.exit(main())
