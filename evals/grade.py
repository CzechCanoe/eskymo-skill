# -*- coding: utf-8 -*-
"""
grade — ohodnotí běhy evalů (evals/evals.json) strojově ověřitelnými tvrzeními.

    python evals/grade.py <iteration-dir>

Očekávaná struktura (skill-creator): <iteration-dir>/eval-<id>-<name>/<config>/run-1/outputs/…
Každému běhu zapíše grading.json (pole expectations: text, passed, evidence).
Ověřuje proti anonymním fixtures z tests/fixtures (stejná data jako inputs evalů).
"""
from __future__ import annotations

import glob
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(ROOT, 'skills', 'csk-eskymo', 'scripts')
FIX = os.path.join(ROOT, 'tests', 'fixtures')
sys.path.insert(0, SCRIPTS)

from domena import vt_skupina  # noqa: E402
from eskymo_ods import Workbook, norm_rgc  # noqa: E402
from registr import Registr  # noqa: E402


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()


def vystupni_ods(outputs):
    c = [f for f in glob.glob(os.path.join(outputs, '**', '*.ods'), recursive=True)
         if 'prepocet' not in os.path.basename(f) and 'recalc' not in f.lower()]
    c.sort(key=lambda f: ('vysled' not in f and 'startov' not in f, -os.path.getmtime(f)))
    return c[0] if c else None


def response(outputs):
    for n in ('response.md', 'odpoved.md'):
        p = os.path.join(outputs, n)
        if os.path.exists(p):
            return open(p, encoding='utf-8', errors='replace').read()
    return ''


def verify(ods, *args):
    r = subprocess.run([sys.executable, os.path.join(SCRIPTS, 'verify_workbook.py'), ods, *args],
                       capture_output=True, text=True, encoding='utf-8', env={**os.environ, 'PYTHONIOENCODING': 'utf-8'})
    return r.stdout


def sl_rows(wb, kat, hlidky=False):
    """[(stc, klic)] ve fyzickém pořadí."""
    out = []
    for r in wb.sheet(f'{kat}_sl').values(max_cols=16)[2:]:
        r = r + [None] * 16
        k = r[15] if hlidky else r[2]
        if k not in (None, '', ' ', 0.0):
            out.append((int(r[1]) if isinstance(r[1], float) else r[1], norm_rgc(k) if not hlidky else k))
    return out


def E(text, passed, evidence):
    return {'text': text, 'passed': bool(passed), 'evidence': str(evidence)[:400]}


# ---------- eval 1 ----------
def grade_vpz(outputs, inputs):
    ex = []
    ods = vystupni_ods(outputs)
    ex.append(E('Výstupní sešit .ods je v outputs/ a vstupní slalom.ods zůstal nezměněný',
                ods and sha(os.path.join(inputs, 'slalom.ods')) == sha(os.path.join(FIX, 'slalom.ods')), ods))
    resp = response(outputs)
    if not ods:
        return ex + [E(t, False, 'chybí výstup') for t in (
            'Startovky obsahují přesně platné lodě závodu 135', 'Dres 13 není použit a čísla jsou unikátní',
            'Každá kategorie začíná číslem končícím 1 (P 2.18.03)', 'Bloky čísel jdou v zadaném pořadí kategorií',
            'VT skupiny v kategorii od nejslabší po nejlepší', 'Po přepočtu se dotáhla všechna jména a oddíly')]
    wb = Workbook(ods)
    # očekávané lodě
    import prihlasky
    ref = prihlasky.over(prihlasky.nacti_csv(os.path.join(FIX, 'prihlasky.csv')), Workbook(os.path.join(FIX, 'slalom.ods')),
                         zavod='135')
    exp = {}
    for l in ref['lode']:
        if not l['kontrola']['problemy']:
            exp.setdefault(l['kat'], []).append(' '.join(sorted(l['rgc'])))
    got = {}
    for kat in ('k1m', 'c1z', 'c2m', 'pzk', 'c1m', 'pzc', 'k1z', 'c2x'):
        rows = sl_rows(wb, kat)
        if rows:
            got[kat] = rows
    got_keys = {k: sorted(' '.join(sorted(str(r[1]).split())) for r in v) for k, v in got.items()}
    exp_keys = {k: sorted(v) for k, v in exp.items()}
    ex.append(E('Startovky obsahují přesně platné lodě závodu 135 (bez RGC 999999, bez lodi jen na závod 136, předžáci v pzk/pzc)',
                got_keys == exp_keys, f'čekáno {dict((k, len(v)) for k, v in exp_keys.items())}, '
                                     f'máme {dict((k, len(v)) for k, v in got_keys.items())}'))
    vse = [s for v in got.values() for s, _ in v]
    ex.append(E('Dres 13 není použit a čísla jsou unikátní', 13 not in vse and len(vse) == len(set(vse)) and all(
        isinstance(s, int) for s in vse), f'{sorted(x for x in vse if isinstance(x, int))[:40]}'))
    ex.append(E('Každá kategorie začíná číslem končícím 1 (P 2.18.03)',
                all(min(s for s, _ in v) % 10 == 1 for v in got.values()),
                {k: min(s for s, _ in v) for k, v in got.items()}))
    order = [k for k in ('k1m', 'c1z', 'c2m', 'pzk', 'c1m', 'pzc', 'k1z', 'c2x') if k in got]
    firsts = [min(s for s, _ in got[k]) for k in order]
    ex.append(E('Bloky čísel jdou v zadaném pořadí kategorií (K1M, C1Ž, C2M, PŽK, C1M, PŽC, K1Ž, C2X)',
                firsts == sorted(firsts), dict(zip(order, firsts))))
    reg = Registr.ze_sesitu(Workbook(os.path.join(FIX, 'slalom.ods')))
    mono = True
    det = {}
    for k, v in got.items():
        grp = [vt_skupina(reg.vt_lode(str(r).split(), k, 'slalom')) for _, r in v]
        det[k] = grp
        if grp != sorted(grp, reverse=True):
            mono = False
    ex.append(E('VT skupiny jdou v kategorii od nejslabší po nejlepší (jako losování Eskyma)', mono,
                {k: g[:12] for k, g in det.items()}))
    v = verify(ods, '--cisla-v-kategorii')
    ex.append(E('Po přepočtu se dotáhla všechna jména a oddíly (žádné #N/A, prázdné jméno)',
                'OK — žádné problémy' in v or ('prázdné jméno' not in v and 'chyba vzorce' not in v and 'PROBLÉMY' not in v),
                v[-300:]))
    ex.append(E('Odpověď pořadateli hlásí neznámé RGC 999999', '999999' in resp, resp[:200]))
    ex.append(E('Odpověď zmiňuje loď přihlášenou jen na jiný závod (136 / neděle)',
                re.search(r'136|neděl', resp, re.I) is not None, ''))
    ex.append(E('Odpověď uvádí další kroky v Eskymu (přepočet Ctrl+Shift+F9 / tisk)',
                re.search(r'Ctrl\+Shift\+F9|přepoč|tisk', resp, re.I) is not None, ''))
    return ex


# ---------- eval 2 ----------
def grade_hlidky(outputs, inputs):
    ex = []
    ods = vystupni_ods(outputs)
    ex.append(E('Výstupní sešit .ods je v outputs/ a vstupní šablona zůstala nezměněná',
                ods and sha(os.path.join(inputs, 'hlidky_sprint.ods')) == sha(os.path.join(FIX, 'hlidky_sprint.ods')), ods))
    resp = response(outputs)
    hl = json.load(open(os.path.join(FIX, 'hlidky.json'), encoding='utf-8'))
    if not ods:
        return ex + [E('výstup', False, 'chybí')]
    wb = Workbook(ods)
    hv = wb.sheet('hlidky').values(max_cols=6)[1:]
    by_hrgc = {r[0]: r for r in hv if r and r[0]}
    ex.append(E('Všech 9 hlídek je v listu hlidky', len(by_hrgc) == 9, f'{len(by_hrgc)} hlídek'))
    import openpyxl
    loni = {}
    lw = openpyxl.load_workbook(os.path.join(FIX, 'loni_hlidky.xlsx'), data_only=True)
    for kat in ('c1m', 'k1z', 'k1m'):
        loni[kat] = [[norm_rgc(x) for x in str(r[4]).split('\n')] for r in list(lw[kat].iter_rows(values_only=True))[2:]]
    nove_ok, obracene_ok, det = True, True, {}
    for kat in ('c1m', 'k1z', 'k1m'):
        rows = sl_rows(wb, kat, hlidky=True)
        prvni = [norm_rgc((by_hrgc.get(h) or [None] * 6)[2]) for _, h in rows]
        det[kat] = prvni
        nova = [h['lode'][0] for h in hl if h['kat'] == kat and h['zdroj'] == 'test-nova'][0]
        if not prvni or prvni[0] != nova:
            nove_ok = False
        ranks = []
        for r in prvni[1:]:
            rk = next((i for i, t in enumerate(loni[kat]) if r in t), None)
            ranks.append(rk)
        if ranks != sorted(ranks, reverse=True) or None in ranks:
            obracene_ok = False
    ex.append(E('V každé kategorii startuje první nová hlídka bez loňské návaznosti', nove_ok, det))
    ex.append(E('Navazující hlídky jdou v obráceném pořadí loňských výsledků (vítěz poslední)', obracene_ok, det))
    stc = {k: [s for s, _ in sl_rows(wb, k, True)] for k in ('c1m', 'k1z', 'k1m')}
    exp = {'c1m': [40, 39, 38], 'k1z': [34, 32, 31], 'k1m': [27, 26, 25]}
    ex.append(E('Čísla 40 dolů, 3 rezervní mezi kategoriemi, 33 vynecháno (C1M 40–38, K1Ž 34/32/31, K1M 27–25)',
                stc == exp, stc))
    fs = ''.join(v for row in wb.sheet('hlidky').values(max_rows=15, max_cols=48, formulas=True) for v in row
                 if isinstance(v, str))
    ex.append(E('Překlep šablony uuper( je ve výstupu opravený', 'uuper(' not in fs, 'uuper(' in fs))
    v = verify(ods)
    ex.append(E('Po přepočtu bez chyb vzorců a s dotaženými oddíly', 'PROBLÉMY' not in v, v[-300:]))
    ex.append(E('Odpověď zmiňuje opravu překlepu ve vzorci šablony', re.search(r'uuper|UPPER|překlep', resp, re.I) is not None, ''))
    ex.append(E('Odpověď upozorňuje na neúplné posádky (zatím jen první lodě) k doplnění',
                re.search(r'doplnit|doplň|první lod|neúpln|RGC2|posádk', resp, re.I) is not None, ''))
    return ex


# ---------- eval 3 ----------
def grade_canoe(outputs, inputs):
    ex = []
    ods = vystupni_ods(outputs)
    ex.append(E('Výstupní sešit .ods je v outputs/ a vstupní slalom.ods zůstal nezměněný',
                ods and sha(os.path.join(inputs, 'slalom.ods')) == sha(os.path.join(FIX, 'slalom.ods')), ods))
    resp = response(outputs)
    if not ods:
        return ex + [E('výstup', False, 'chybí')]
    wb = Workbook(ods)
    k1z = sl_rows(wb, 'k1z')
    ex.append(E('Všech 6 žen z třídy K1Z (české kódy) je ve startovce k1z', len(k1z) == 6, len(k1z)))
    ex.append(E('Odpověď hlásí neznámou třídu XYZ (závodník nebyl převeden)', 'XYZ' in resp, ''))
    v = verify(ods, '--vysledky', '--cisla-v-kategorii')
    ex.append(E('Každá jízda zapsaných závodníků má čas nebo stav', 'bez času i stavu' not in v and 'nezadána' not in v,
                v[-300:]))
    ok999, det = True, []
    for kat in ('k1m', 'k1z', 'c1m', 'c2m'):
        for r in wb.sheet(kat).values(max_cols=16)[2:]:
            r = r + [None] * 16
            for tc, pc in ((11, 12), (14, 15)):
                if isinstance(r[tc], str) and r[tc].strip().upper().startswith(('DNS', 'DNF', 'DSQ')):
                    det.append((kat, r[tc], r[pc]))
                    if r[pc] != 999:
                        ok999 = False
    ex.append(E('Stavy DNS/DNF mají v penalizaci 999', ok999 and det, det[:6]))
    c1m = [k for _, k in sl_rows(wb, 'c1m')]
    cizi = [r[0] for r in wb.sheet('cizi').values(max_cols=2)[1:] if r and r[0]]
    ex.append(E('Cizinec SMITH spárován s existujícím A90002 (bez duplicitního A-kódu)',
                'A90002' in c1m and len(cizi) == 2, f'c1m={c1m[-3:]}, cizi={cizi}'))
    c2 = [k for _, k in sl_rows(wb, 'c2m')]
    ex.append(E('Obě deblové posádky mají dvě RGC (i starý slepený formát rozdělen)',
                len(c2) == 2 and all(len(str(x).split()) == 2 for x in c2), c2))
    from inspect_workbook import read_param
    p = {k: v['hodnota'] for k, v in read_param(wb)['pole'].items()}
    ex.append(E('param: datum 15.08.26 a číslo závodu 135', str(p.get('Datum závodu')) == '15.08.26' and
                p.get('Číslo závodu') in (135, 135.0), (p.get('Datum závodu'), p.get('Číslo závodu'))))
    ex.append(E('Odpověď připomíná přepočet (Ctrl+Shift+F9) a body tlačítkem v Eskymu',
                re.search(r'Ctrl\+Shift\+F9|přepoč', resp, re.I) is not None and re.search(r'bod', resp, re.I) is not None, ''))
    return ex


# ---------- eval 4 ----------
def grade_cpw(outputs, inputs):
    resp = response(outputs)
    t = resp.lower()
    return [
        E('Obrácené průběžné pořadí ČPw (od nejhorších k nejlepším)', re.search(r'od nejhorš|obrácen', t), ''),
        E('Průběžné pořadí se bere jen z klasických sjezdů', re.search(r'klasick', t) and re.search(r'jen|pouze', t), ''),
        E('1. a 2. ČPw podle loňského žebříčku', re.search(r'1\. a 2\.|prvn\w* (a|dv)\w* .*lo[ňn]', t) and 'lo' in t, ''),
        E('Čísla od nejvyššího po nejnižší, poslední v kategorii má číslo 1',
          re.search(r'nejvyšší', t) and re.search(r'(číslo|č\.)\s*1\b|jedničk', t), ''),
        E('Podklady (průběžné pořadí) dodá počtářka žebříčku', re.search(r'počtář', t), ''),
        E('Nezařazené Směrnice pro ČPw výslovně neřeší → rozhodne pořadatel/VR (bez vymyšleného pravidla)',
          re.search(r'neřeš|neuvád|neupravuj|není (výslovně )?(stanoven|uveden|upraven)|mlč', t), ''),
        E('Cituje Směrnice pro závodění 2026', re.search(r'směrnic\w* .{0,40}2026|s26', t), ''),
        E('Uvádí pořadí kategorií C1M, K1Ž, K1M, C1Ž, C2M nebo minimální interval', re.search(
            r'c1m,?\s*k1[žz],?\s*k1m|30\s*s', t), ''),
    ]


GRADERS = {'vpz-startovka': grade_vpz, 'mcr-hlidky': grade_hlidky, 'canoe123-vysledky': grade_canoe,
           'pravidla-cpw': grade_cpw}


def main(it):
    for ev in sorted(glob.glob(os.path.join(it, 'eval-*'))):
        name = re.sub(r'^eval-\d+-', '', os.path.basename(ev))
        g = GRADERS[name]
        for run in sorted(glob.glob(os.path.join(ev, '*', 'run-*'))):
            outputs = os.path.join(run, 'outputs')
            try:
                ex = g(outputs, os.path.join(ev, 'inputs'))
            except Exception as e:  # hodnocení nesmí spadnout kvůli rozbitému výstupu
                ex = [E('Hodnocení proběhlo', False, repr(e))]
            passed = sum(1 for e in ex if e['passed'])
            gr = {'expectations': ex, 'summary': {'passed': passed, 'failed': len(ex) - passed, 'total': len(ex),
                                                  'pass_rate': round(passed / len(ex), 3) if ex else 0}}
            tf = os.path.join(run, 'timing.json')
            if os.path.exists(tf):
                gr['timing'] = json.load(open(tf, encoding='utf-8'))
            json.dump(gr, open(os.path.join(run, 'grading.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
            print(f'{os.path.basename(ev)} {os.path.basename(os.path.dirname(run))}: {passed}/{len(ex)}')


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    main(sys.argv[1])
